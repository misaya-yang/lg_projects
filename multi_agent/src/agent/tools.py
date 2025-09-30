"""Utility tools available to the research agent."""

from __future__ import annotations

import asyncio
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, List, Optional, cast

from docx import Document
from langchain_tavily import TavilySearch  # type: ignore[import-not-found]

from src.agent.configuration import Configuration


async def search(query: str) -> Optional[dict[str, Any]]:
    """Search for general web results."""

    configuration = Configuration.from_context()
    wrapped = TavilySearch(max_results=configuration.max_search_results)
    return cast(dict[str, Any], await wrapped.ainvoke({"query": query}))


async def add(a: int, b: int) -> int:
    """Add two numbers."""

    return a + b


async def sub(a: int, b: int) -> int:
    """Subtract two numbers."""

    return a - b


async def mul(a: int, b: int) -> int:
    """Multiply two numbers."""

    return a * b


async def div(a: int, b: int) -> float:
    """Divide two numbers."""

    return a / b


def _markdown_to_docx(markdown_text: str, output_path: Path, title: str | None = None) -> str:
    document = Document()
    if title:
        document.add_heading(title, level=0)

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if not stripped:
            document.add_paragraph()
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            document.add_heading(text, level=min(level, 9))
            continue
        if stripped.startswith(('-', '*')):
            document.add_paragraph(stripped[1:].strip(), style="List Bullet")
            continue
        if stripped[:2].isdigit() and stripped[2:3] == ".":
            _, _, text = stripped.partition(".")
            document.add_paragraph(text.strip(), style="List Number")
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph(stripped[1:].strip())
            paragraph.style = "Intense Quote"
            continue
        document.add_paragraph(line)

    document.save(output_path)
    return str(output_path)


async def markdown_to_docx(markdown_text: str, output_path: str | None = None, title: str | None = None) -> str:
    """Convert markdown content into a `.docx` document and return the path."""

    configuration = Configuration.from_context()
    provided_path = Path(output_path) if output_path else None
    if provided_path and provided_path.suffix == ".docx":
        target_path = provided_path
    else:
        base_dir = provided_path or Path(configuration.default_docx_directory)
        base_dir.mkdir(parents=True, exist_ok=True)
        filename = f"deep_research_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.docx"
        target_path = base_dir / filename

    target_path.parent.mkdir(parents=True, exist_ok=True)
    return await asyncio.to_thread(_markdown_to_docx, markdown_text, target_path, title)


TOOLS: List[Callable[..., Any]] = [search, add, sub, mul, div, markdown_to_docx]
